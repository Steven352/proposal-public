import copy
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from proposal_app import github_library


class MemoryLibrary:
    def __init__(self):
        self.files = {}
        self.values = {}

    def read_json(self, path, default):
        return copy.deepcopy(self.values.get(path, default))

    def write_bytes(self, path, data, message):
        self.files[path] = data

    def update_json(self, path, default, updater, message):
        self.values[path] = updater(copy.deepcopy(self.values.get(path, default)))
        return self.values[path]


def proposal_bytes(number='P026-200'):
    doc = Document()
    doc.add_paragraph(number)
    doc.add_heading('Introduction', 1)
    doc.add_paragraph('Geotechnical assessment for a proposed building.')
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


class AutomaticLibraryTest(unittest.TestCase):
    def test_saves_original_and_search_index_without_draft_or_learning(self):
        self.assertTrue(hasattr(github_library, 'save_pdf_source_to_library'))
        client = MemoryLibrary()
        data = proposal_bytes()
        with tempfile.TemporaryDirectory() as folder, patch.object(
            github_library, 'HISTORICAL_DIR', Path(folder) / 'proposals'
        ), patch.object(github_library, 'LIBRARY_INDEX_ADDITIONS', Path(folder) / 'index.json'):
            for _ in range(2):
                _, state, path = github_library.save_pdf_source_to_library(client, 'P026-200.docx', data)
            self.assertEqual(client.files[path], data)
            self.assertEqual((Path(folder) / 'proposals/P026-200.docx').read_bytes(), data)
            self.assertEqual(len(client.values[github_library.ADDITIONS_PATH]['proposals']), 1)
            self.assertEqual(len(state['records']), 1)
            self.assertEqual(state['records'][0]['candidates'], [])
            self.assertEqual(state['records'][0]['draft_sha256'], '')

    def test_excluded_proposal_is_rejected_before_any_write(self):
        self.assertTrue(hasattr(github_library, 'save_pdf_source_to_library'))
        client = MemoryLibrary()
        with self.assertRaises(ValueError):
            github_library.save_pdf_source_to_library(client, 'P026-133.docx', proposal_bytes('P026-133'))
        self.assertEqual(client.files, {})

class ExistingReviewTest(unittest.TestCase):
    def test_regenerating_pdf_preserves_prior_review_analysis(self):
        import hashlib
        client = MemoryLibrary()
        data = proposal_bytes()
        record = {'id': hashlib.sha256(data).hexdigest()[:16], 'summary': 'Reviewed edits',
                  'candidates': [{'key': 'wording-rule'}], 'repository_path': 'data/historical_proposals/P026-200.docx'}
        client.values[github_library.STATE_PATH] = {'version': 1, 'records': [record], 'approved_rules': []}
        client.read_json = lambda path, default: copy.deepcopy(client.values.get(path, default))
        with tempfile.TemporaryDirectory() as folder, patch.object(github_library, 'HISTORICAL_DIR', Path(folder)), patch.object(github_library, 'LIBRARY_INDEX_ADDITIONS', Path(folder) / 'index.json'):
            _, state, _ = github_library.save_pdf_source_to_library(client, 'P026-200.docx', data)
        self.assertEqual(state['records'][0]['summary'], 'Reviewed edits')
        self.assertEqual(state['records'][0]['candidates'], [{'key': 'wording-rule'}])
