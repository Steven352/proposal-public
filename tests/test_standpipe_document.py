import io
import unittest
from docx import Document
from proposal_app import document_builder

WORDING = ('Installation of 25 mm diameter standpipe piezometers in all boreholes, backfilled with '
           'cuttings and capped with bentonite. Groundwater readings will be taken at the end of '
           'drilling and during one follow-up visit two weeks after the field program.')


class StandpipeDocumentTest(unittest.TestCase):
    def test_missing_scope_is_added_once_inside_field_program(self):
        self.assertTrue(hasattr(document_builder, 'ensure_standpipe_scope'))
        doc = Document()
        doc.add_heading('Geotechnical Field Program', 1)
        doc.add_paragraph('Drilling and sampling of two boreholes.', style='List Bullet')
        doc.add_heading('Cost of Geotechnical Services', 1)
        doc.add_paragraph('Cost remains unchanged.')
        data = io.BytesIO()
        doc.save(data)
        result = document_builder.ensure_standpipe_scope(data.getvalue())
        result = document_builder.ensure_standpipe_scope(result)
        texts = [p.text for p in Document(io.BytesIO(result)).paragraphs]
        self.assertEqual(texts.count(WORDING), 1)
        self.assertLess(texts.index(WORDING), texts.index('Cost of Geotechnical Services'))
        self.assertIn('Cost remains unchanged.', texts)

    def test_test_pit_only_document_is_unchanged(self):
        self.assertTrue(hasattr(document_builder, 'ensure_standpipe_scope'))
        doc = Document()
        doc.add_heading('Geotechnical Field Program', 1)
        doc.add_paragraph('Excavation of two test pits.')
        doc.add_heading('Cost of Geotechnical Services', 1)
        out = io.BytesIO()
        doc.save(out)
        self.assertEqual(document_builder.ensure_standpipe_scope(out.getvalue()), out.getvalue())

    def test_generated_proposal_uses_approved_wording_when_ai_returns_old_clause(self):
        doc = Document()
        doc.add_heading('Geotechnical Field Program', 1)
        doc.add_paragraph('Installation of standpipe piezometers in selected boreholes.')
        doc.add_heading('Cost of Geotechnical Services', 1)
        out = io.BytesIO()
        doc.save(out)
        result = document_builder.ensure_standpipe_scope(out.getvalue(), required=True)
        texts = [p.text for p in Document(io.BytesIO(result)).paragraphs]
        self.assertIn(WORDING, texts)
        self.assertFalse(any('selected boreholes' in text for text in texts))
