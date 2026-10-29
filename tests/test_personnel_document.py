import io
import tempfile
import unittest
from pathlib import Path

from docx import Document

from proposal_app import document_builder
from proposal_app.models import CostLineItem, DraftContent, ParagraphBlock, ProposalFacts


class PersonnelDocumentTest(unittest.TestCase):
    def template(self, include_personnel=True, scope_heading="Scope of Work and Cost Summary"):
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Current project header"
        document.add_heading("Introduction", 1)
        document.add_paragraph("Introduction content")
        if include_personnel:
            document.add_heading("Proposed Project Personnel", 1)
            document.add_paragraph("Approved personnel biography")
        if scope_heading:
            document.add_heading(scope_heading, 1)
        document.add_paragraph("Scope content")
        return document

    def test_unchecked_removes_entire_section_and_preserves_header(self):
        document = self.template()
        document_builder.configure_personnel_section(document, False)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertNotIn("Proposed Project Personnel", text)
        self.assertNotIn("Approved personnel biography", text)
        self.assertIn("Introduction content", text)
        self.assertIn("Scope content", text)
        self.assertEqual(document.sections[0].header.paragraphs[0].text, "Current project header")

    def test_checked_preserves_existing_template_section(self):
        document = self.template()
        original = document.element.xml
        document_builder.configure_personnel_section(document, True)
        self.assertEqual(document.element.xml, original)

    def test_checked_imports_template_section_if_selected_layout_lacks_it(self):
        document = self.template(False)
        with tempfile.TemporaryDirectory() as directory:
            self.template().save(Path(directory) / "approved-template.docx")
            document_builder.configure_personnel_section(document, True, Path(directory))
        texts = [p.text for p in document.paragraphs]
        self.assertIn("Approved personnel biography", texts)
        self.assertLess(texts.index("Introduction content"), texts.index("Proposed Project Personnel"))
        self.assertLess(texts.index("Approved personnel biography"), texts.index("Scope of Work and Cost Summary"))
        buffer = io.BytesIO()
        document.save(buffer)
        self.assertEqual(Document(buffer).sections[0].header.paragraphs[0].text, "Current project header")

    def test_missing_approved_personnel_template_reports_actionable_error(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "personnel.*template|template.*personnel"):
                document_builder.configure_personnel_section(self.template(False), True, Path(directory))

    def test_checked_imports_before_selected_cost_summary_heading(self):
        document = self.template(False, "Cost Summary")
        with tempfile.TemporaryDirectory() as directory:
            self.template().save(Path(directory) / "approved-template.docx")
            document_builder.configure_personnel_section(document, True, Path(directory))
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("Approved personnel biography", texts)
        self.assertLess(texts.index("Approved personnel biography"), texts.index("Cost Summary"))
        self.assertIn("Scope content", texts)

    def test_unchecked_removes_personnel_subsections_before_cost_summary(self):
        document = self.template(True, "Cost Summary")
        boundary = next(p for p in document.paragraphs if p.text == "Cost Summary")
        boundary.insert_paragraph_before("Personnel qualifications", style="Heading 2")
        boundary.insert_paragraph_before("Personnel subsection content")
        document_builder.configure_personnel_section(document, False)
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertNotIn("Proposed Project Personnel", texts)
        self.assertNotIn("Approved personnel biography", texts)
        self.assertNotIn("Personnel qualifications", texts)
        self.assertNotIn("Personnel subsection content", texts)
        self.assertIn("Introduction content", texts)
        self.assertIn("Cost Summary", texts)
        self.assertIn("Scope content", texts)

    def test_checked_accepts_approved_donor_with_alternate_scope_heading(self):
        document = self.template(False)
        with tempfile.TemporaryDirectory() as directory:
            self.template(True, "Cost Summary").save(Path(directory) / "approved-template.docx")
            document_builder.configure_personnel_section(document, True, Path(directory))
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("Approved personnel biography", texts)
        self.assertNotIn("Cost Summary", texts)
        self.assertLess(
            texts.index("Approved personnel biography"),
            texts.index("Scope of Work and Cost Summary"),
        )

    def test_personnel_option_preserves_field_program_without_scope_heading(self):
        for include in (True, False):
            with self.subTest(include=include):
                document = self.template(not include, None)
                document.add_heading("Geotechnical Field Program", 3)
                document.add_paragraph("Field program content")
                with tempfile.TemporaryDirectory() as directory:
                    self.template().save(Path(directory) / "approved-template.docx")
                    document_builder.configure_personnel_section(document, include, Path(directory))
                texts = [paragraph.text for paragraph in document.paragraphs]
                self.assertEqual("Proposed Project Personnel" in texts, include)
                self.assertEqual("Approved personnel biography" in texts, include)
                self.assertIn("Introduction content", texts)
                self.assertIn("Geotechnical Field Program", texts)
                self.assertIn("Field program content", texts)

    def test_build_docx_preserves_alternate_scope_with_either_personnel_option(self):
        facts = ProposalFacts(
            proposal_number="P026-999", client_name="Example Client",
            project_name="Example Assessment", project_location="Example Location",
            cost_items=[CostLineItem(
                section="Field Program", item="Field services", unit="hr",
                estimate=2, rate=100, total=200,
            )],
        )
        draft = DraftContent(
            introduction="Updated introduction for the example assessment.",
            field_program_intro="Updated field program introduction.",
            field_program_paragraphs=[ParagraphBlock(text="Excavation of two test pits.")],
            cost_intro="Updated estimate introduction.",
            terms_and_conditions="Terms remain controlled by the template.",
            closure_paragraph_1="Example closure one.",
            closure_paragraph_2="Example closure two.",
            work_authorization_scope="Example assessment", email_subject="Example proposal",
            email_body="Example proposal email.",
        )
        for include in (True, False):
            with self.subTest(include=include), tempfile.TemporaryDirectory() as directory:
                document = Document()
                for text in (
                    "October 1, 2026\tAlmor Proposal No.: P026-001", "Original Client", "",
                    "Attention: Original Contact", "Re:\tOriginal Assessment", "Original Location",
                ):
                    document.add_paragraph(text)
                document.add_heading("Introduction", 1)
                document.add_paragraph("Original introduction.")
                document.add_heading("Proposed Project Personnel", 1)
                document.add_paragraph("Approved personnel biography")
                document.add_heading("Cost Summary", 1)
                document.add_heading("Geotechnical Field Program", 3)
                document.add_paragraph("Original field program.")
                document.add_heading("Cost of Geotechnical Services", 3)
                document.add_paragraph("Original estimate introduction.")
                document.add_table(rows=5, cols=6)
                document.add_paragraph("The estimated fee is $9.00, excluding taxes.")
                document.add_heading("Terms and Conditions", 1)
                document.add_paragraph("Original approved terms.")
                document.add_heading("Closure", 1)
                for text in (
                    "Closure one.", "Closure two.", "Respectfully Submitted,", "Example Company",
                    "Prepared by\tReviewed By", "Example Author\tExample Reviewer",
                    "Example title\tExample title", "Original document identifier",
                ):
                    document.add_paragraph(text)
                template = Path(directory) / "example-template.docx"
                document.save(template)
                data, _ = document_builder.build_docx(template, facts, draft, include)
                result = Document(io.BytesIO(data))
                texts = [paragraph.text for paragraph in result.paragraphs]
                self.assertIn(draft.introduction, texts)
                self.assertIn("Cost Summary", texts)
                self.assertIn("Geotechnical Field Program", texts)
                self.assertIn(draft.field_program_intro, texts)
                self.assertIn("Excavation of two test pits.", texts)
                self.assertIn("The estimated fee is $200.00, excluding taxes.", texts)
                self.assertIn("Original approved terms.", texts)
                self.assertEqual("Proposed Project Personnel" in texts, include)
                self.assertEqual("Approved personnel biography" in texts, include)
                table_text = "\n".join(cell.text for row in result.tables[0].rows for cell in row.cells)
                self.assertIn("$200.00", table_text)
