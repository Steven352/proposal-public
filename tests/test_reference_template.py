import io
import unittest
from pathlib import Path

from docx import Document

from proposal_app.config import STANDARD_TEMPLATE_PATH
from proposal_app import document_builder
from tests.test_standard_document import sample_draft, sample_facts


class ReferenceTemplateTest(unittest.TestCase):
    def test_standard_template_is_packaged_and_has_reference_page_system(self):
        self.assertTrue(STANDARD_TEMPLATE_PATH.is_file(), STANDARD_TEMPLATE_PATH)
        document = Document(STANDARD_TEMPLATE_PATH)
        section = document.sections[0]
        self.assertTrue(section.different_first_page_header_footer)
        self.assertAlmostEqual(section.page_width.inches, 8.5, places=2)
        self.assertAlmostEqual(section.page_height.inches, 11.0, places=2)
        body = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Scope of Work and Cost Summary", body)
        self.assertIn("Terms and Conditions", body)
        self.assertNotIn("Ashley Robin Ross", body)
        self.assertNotIn("P026-187", body)

    def test_generation_preserves_reference_front_matter_slots(self):
        data, _ = document_builder.build_standard_docx(
            sample_facts(), sample_draft(), include_proposed_project_personnel=False
        )
        document = Document(io.BytesIO(data))
        body = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Attention: Alex Example", body)
        self.assertIn("Almor Proposal No.: P026-999", body)
        self.assertIn("Current Example Development", body)
        self.assertIn("Example County, Alberta", body)
        self.assertNotIn("Client Contact", body)
        first_header = " ".join(document.sections[0].first_page_header._element.xpath(".//w:t/text()"))
        self.assertNotIn("Project Name", first_header)
        repeating = document.sections[0].header
        self.assertEqual(repeating.paragraphs[0].style.name, "Header Client name")
        self.assertEqual(repeating.paragraphs[1].style.name, "Header Body")
        self.assertEqual(repeating.paragraphs[0].runs[0].font.size.pt, 9)
        self.assertEqual(repeating.paragraphs[1].runs[0].font.size.pt, 8)
        self.assertEqual(str(repeating.paragraphs[1].runs[0].font.color.rgb), "663300")


if __name__ == "__main__":
    unittest.main()
