import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from streamlit.testing.v1 import AppTest
from docx import Document

from proposal_app import document_builder
from proposal_app.models import ProposalFacts, ReferenceMatch
from test_standard_document import make_standard_templates, sample_draft


class ProposalNumberUITest(unittest.TestCase):
    def setUp(self):
        for target, kwargs in (
            ("proposal_app.secure_bundle.ensure_private_assets", {}),
            ("proposal_app.knowledge.load_index", {"return_value": []}),
        ):
            mocked = patch(target, **kwargs)
            mocked.start()
            self.addCleanup(mocked.stop)

    def app(self, legacy_state=None):
        app = AppTest.from_file(
            str(Path(__file__).resolve().parents[1] / "streamlit_app.py"),
            default_timeout=20,
        )
        app.secrets["OPENAI_API_KEY"] = "test-key"
        app.secrets["GITHUB_LIBRARY_TOKEN"] = ""
        app.session_state["access_granted"] = True
        app.session_state["access_role"] = "standard"
        for key, value in (legacy_state or {}).items():
            app.session_state[key] = value
        app.run()
        self.assertFalse(app.exception)
        return app

    def number_fields(self, app):
        return [field for field in app.text_input if field.label == "Proposal number"]

    def test_restore_draft_updates_existing_widgets_and_editor(self):
        import streamlit as st
        from proposal_app.draft_state import DRAFT_SCHEMA
        payload = {"schema": DRAFT_SCHEMA, "version": 1,
                   "facts": {"proposal_number": "P026-888", "client_name": "Restored Client"},
                   "request_email": "Saved request", "notes": "Saved notes",
                   "include_proposed_project_personnel": False}
        original = st.file_uploader
        def upload(label, *args, **kwargs):
            if label == "Upload an Almor draft":
                return io.BytesIO(json.dumps(payload).encode())
            return original(label, *args, **kwargs)
        with patch("streamlit.file_uploader", side_effect=upload):
            app = self.app()
            self.click(app, "Enter manually")
            self.click(app, "Restore draft")
            self.assertFalse(app.error)
            self.assertEqual(self.number_fields(app)[0].value, "P026-888")
            self.assertEqual(app.text_input(key="edit_client_name").value, "Restored Client")
            self.assertEqual(app.text_area(key="request_email").value, "Saved request")
            self.assertFalse(app.checkbox(key="include_proposed_project_personnel").value)

    def click(self, app, label):
        next(button for button in app.button if button.label == label).click().run()
        self.assertFalse(app.exception)

    def test_manual_entry_asks_for_proposal_number_once_and_preserves_edits(self):
        app = self.app()
        self.number_fields(app)[0].set_value("P026-240").run()
        self.click(app, "Enter manually")
        self.assertEqual(len(self.number_fields(app)), 1)
        self.assertEqual(self.number_fields(app)[0].value, "P026-240")
        self.assertEqual(app.session_state["proposal_facts"]["proposal_number"], "P026-240")

        self.number_fields(app)[0].set_value("P026-241").run()
        self.assertEqual(app.session_state["proposal_facts"]["proposal_number"], "P026-241")
        self.click(app, "Enter manually")
        self.assertEqual(self.number_fields(app)[0].value, "P026-241")
        self.assertEqual(app.session_state["proposal_facts"]["proposal_number"], "P026-241")

    def test_analysis_keeps_user_number_even_if_extraction_returns_a_different_number(self):
        app = self.app()
        self.number_fields(app)[0].set_value("P026-242").run()
        next(field for field in app.text_area if field.label == "Request email").set_value(
            "Please prepare an assessment proposal for Example Client."
        ).run()
        with patch(
            "proposal_app.ai.ProposalAI.extract_facts",
            return_value=ProposalFacts(proposal_number="P026-001", client_name="Example Client"),
        ):
            self.click(app, "Analyze request")
        self.assertEqual(app.session_state["proposal_facts"]["proposal_number"], "P026-242")
        self.assertEqual(len(self.number_fields(app)), 1)
        self.assertEqual(self.number_fields(app)[0].value, "P026-242")

    def test_blank_number_requires_the_single_entry_before_generation(self):
        app = self.app()
        self.click(app, "Enter manually")
        self.assertEqual(len(self.number_fields(app)), 1)
        generate = next(button for button in app.button if button.label == "Generate Word proposal")
        self.assertTrue(generate.disabled)
        self.assertTrue(any("proposal number" in warning.value.lower() for warning in app.warning))
        self.number_fields(app)[0].set_value("P026-247").run()
        generate = next(button for button in app.button if button.label == "Generate Word proposal")
        self.assertFalse(generate.disabled)
        self.assertEqual(app.session_state["proposal_facts"]["proposal_number"], "P026-247")

    def test_number_change_removes_previous_word_download(self):
        app = self.app()
        self.number_fields(app)[0].set_value("P026-243").run()
        self.click(app, "Enter manually")
        app.session_state["generated_outputs"] = {
            "docx": b"previous word", "docx_name": "P026-243 Example.docx",
            "template": "Standard Almor proposal template",
        }
        app.run()
        self.assertTrue(any(item.label == "Download Word proposal" for item in app.get("download_button")))
        self.number_fields(app)[0].set_value("P026-244").run()
        self.assertNotIn("generated_outputs", app.session_state)
        self.assertEqual(app.session_state["proposal_facts"]["proposal_number"], "P026-244")
        self.assertFalse(any(item.label == "Download Word proposal" for item in app.get("download_button")))

    def test_legacy_layout_download_is_removed_but_standard_download_is_retained(self):
        for template, retained in (
            ("Historic project layout.docx", False),
            ("Standard Almor proposal template", True),
        ):
            with self.subTest(template=template):
                app = self.app({
                    "proposal_number": "P026-248",
                    "proposal_facts": ProposalFacts(proposal_number="P026-248").model_dump(mode="json"),
                    "generated_outputs": {
                        "docx": b"previous word", "docx_name": "P026-248 Example.docx",
                        "template": template,
                    },
                })
                self.assertEqual("generated_outputs" in app.session_state, retained)
                self.assertEqual(
                    any(item.label == "Download Word proposal" for item in app.get("download_button")),
                    retained,
                )

    def test_number_migration_uses_stored_facts_when_legacy_editor_is_blank(self):
        for legacy_number, expected in (
            ("", "P026-249"), ("   ", "P026-249"), ("P026-250", "P026-250"),
        ):
            with self.subTest(legacy_number=legacy_number):
                app = self.app({
                    "edit_proposal_number": legacy_number,
                    "proposal_facts": ProposalFacts(proposal_number="P026-249").model_dump(mode="json"),
                })
                self.assertEqual(len(self.number_fields(app)), 1)
                self.assertEqual(self.number_fields(app)[0].value, expected)
                self.assertEqual(app.session_state["proposal_facts"]["proposal_number"], expected)

    def test_generation_uses_current_number_and_standard_layout_with_content_references(self):
        references = [
            ReferenceMatch(filename="Ranked content reference.docx", project_type="Example", methods=[],
                           score=10, has_cost_table=False, sections={}),
            ReferenceMatch(filename="Different historical layout.docx", project_type="Example", methods=[],
                           score=5, has_cost_table=True, sections={}),
        ]
        for matches in (references, []):
            with self.subTest(has_references=bool(matches)), tempfile.TemporaryDirectory() as directory:
                base, _ = make_standard_templates(Path(directory))

                def draft_content(_client, facts, content_references, rules_path):
                    introduction = (
                        "Wording based on Ranked content reference.docx."
                        if content_references and content_references[0].filename == "Ranked content reference.docx"
                        else "Wording without historical references."
                    )
                    return sample_draft().model_copy(update={"introduction": introduction})

                with patch.object(document_builder, "STANDARD_TEMPLATE_PATH", base, create=True), \
                     patch("proposal_app.knowledge.retrieve_references", return_value=matches), \
                     patch("proposal_app.ai.ProposalAI.draft_content", autospec=True, side_effect=draft_content):
                    app = self.app()
                    self.number_fields(app)[0].set_value("P026-245").run()
                    self.click(app, "Enter manually")
                    frame = app.session_state["cost_frame"].copy()
                    frame.loc[0, "Est."] = 2.0
                    app.session_state["cost_frame"] = frame
                    app.run()
                    self.number_fields(app)[0].set_value("P026-246").run()
                    self.click(app, "Generate Word proposal")
                self.assertFalse(app.error)
                output = app.session_state["generated_outputs"]
                self.assertTrue(output["docx_name"].startswith("P026-246"))
                self.assertEqual(output["template"], "Standard Almor proposal template")
                document = Document(io.BytesIO(output["docx"]))
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                self.assertIn("P026-246", text)
                self.assertNotIn("P026-245", text)
                self.assertIn(
                    "Wording based on Ranked content reference.docx." if matches
                    else "Wording without historical references.",
                    text,
                )
                self.assertEqual(
                    [cell.text for cell in document.tables[0].rows[0].cells],
                    ["Item", "Description", "Unit", "Est.", "Rate", "Total"],
                )


if __name__ == "__main__":
    unittest.main()
