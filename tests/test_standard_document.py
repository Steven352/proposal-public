import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, Twips

from proposal_app import document_builder
from proposal_app.models import CostLineItem, DraftContent, ParagraphBlock, ProposalFacts


def sample_facts():
    return ProposalFacts(
        proposal_number="P026-999", proposal_date="2026-10-29",
        client_name="Current Example Client Ltd.", client_address="100 Example Avenue",
        contact_name="Alex Example", project_name="Current Example Development",
        project_location="Example County, Alberta", investigation_methods=["test pits"],
        cost_items=[CostLineItem(
            section="Field Program", item="Excavator", description="Field investigation",
            unit="hr", estimate=3, rate=150, total=450,
        )],
    )


def sample_draft():
    return DraftContent(
        introduction="Current example introduction.",
        field_program_intro="Current example field program.",
        field_program_paragraphs=[ParagraphBlock(text="Excavate three test pits.", style="list")],
        cost_intro="The following estimate applies to this scope.",
        terms_and_conditions="AI text must not replace the approved terms.",
        closure_paragraph_1="AI text must not replace the approved closure.",
        closure_paragraph_2="AI alternate closure.", work_authorization_scope="Example investigation",
        email_subject="Example proposal", email_body="Example email.",
    )


def _table(document, standard):
    table = document.add_table(rows=5, cols=6)
    table.style = "Table Grid"
    widths = [1.1, 2.5, .4, .5, .75, .85] if standard else [1] * 6
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    headings = ["Item", "Description", "Unit", "Est.", "Rate", "Total"]
    for index, cell in enumerate(table.rows[0].cells):
        cell.text = headings[index] if standard else f"Unwanted layout {index}"
    for row in table.rows:
        for cell in row.cells:
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(3 if standard else 12)
            paragraph.paragraph_format.line_spacing = 1.0
            if not paragraph.runs:
                paragraph.add_run("")
            paragraph.runs[0].font.name = "Arial"
            paragraph.runs[0].font.size = Pt(8 if standard else 11)
    return table


def make_standard_templates(directory):
    """Generic fixed document plus a decoy cost file that generation does not require."""
    directory = Path(directory)
    document = Document()
    document.styles.add_style("Body", WD_STYLE_TYPE.PARAGRAPH)
    document.styles["Body"].font.name = "Arial"
    document.styles["Body"].font.size = Pt(11)
    document.sections[0].different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True
    for header in (document.sections[0].header, document.sections[0].first_page_header):
        for index, text in enumerate((
            "Prior Header Client Alias Inc.", "Proposal for Geotechnical Services",
            "Prior Header Project Alias", "Almor Proposal No.: P026-111",
        )):
            paragraph = header.paragraphs[0] if index == 0 else header.add_paragraph()
            paragraph.text = text
    document.sections[0].even_page_header._element.append(parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict>'
        '<v:shape id="ExampleHeader" style="width:400pt;height:70pt"><v:textbox>'
        '<w:txbxContent>'
        '<w:p><w:r><w:t>Prior Textbox Client Alias Inc.</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Proposal for Geotechnical Services</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Prior Textbox Project Alias</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Almor Proposal No.: P026-222</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
    ))
    for text in (
        "October 1, 2026\tAlmor Proposal No.: P026-001", "Prior Front Client", "",
        "Attention: Prior Contact", "Re:\tPrior Front Project", "Prior Front Location",
    ):
        document.add_paragraph(text)
    document.add_heading("Introduction", 1)
    document.add_paragraph("Prior introduction.", "Body")
    document.add_heading("Proposed Project Personnel", 1)
    document.add_paragraph("Approved personnel biography.", "Body")
    document.add_heading("Scope of Work and Cost Summary", 1)
    document.add_heading("Scope of Work", 2)
    document.add_heading("Geotechnical Field Program", 3)
    document.add_paragraph("Prior field program.", "Body")
    document.add_heading("Cost of Geotechnical Services", 3)
    document.add_paragraph("Approved cost lead-in.", "Body")
    _table(document, False)
    document.add_paragraph("The total estimated services cost is $9.00, excluding GST.", "Body")
    document.add_heading("Terms and Conditions", 1)
    document.add_paragraph("Approved standard terms.", "Body")
    document.add_heading("Closure", 1)
    for text in (
        "Approved closure one.", "Approved closure two.", "Respectfully Submitted,",
        "ALMOR TESTING SERVICES LTD.", "Prepared by\tReviewed By", "Example Author\tExample Reviewer",
        "Example title\tExample title", "Prior document identifier",
    ):
        document.add_paragraph(text, "Body")
    standard_path = directory / "standard-proposal.docx"
    document.save(standard_path)
    cost_document = Document()
    _table(cost_document, True)
    cost_path = directory / "standard-cost-table.docx"
    cost_document.save(cost_path)
    return standard_path, cost_path


class StandardDocumentTest(unittest.TestCase):
    def build(self, base, facts=None, draft=None):
        build = getattr(document_builder, "build_standard_docx", None)
        self.assertIsNotNone(build, "Generation must expose a fixed standard-template entry point")
        with patch.object(document_builder, "STANDARD_TEMPLATE_PATH", base, create=True):
            data, filename = build(facts or sample_facts(), draft or sample_draft())
        return Document(io.BytesIO(data)), filename

    def test_all_project_headers_use_current_facts_despite_historical_header_aliases(self):
        with tempfile.TemporaryDirectory() as directory:
            base, _ = make_standard_templates(directory)
            result, filename = self.build(base)
        self.assertTrue(filename.startswith("P026-999 "))
        for header in (
            result.sections[0].header, result.sections[0].first_page_header,
            result.sections[0].even_page_header,
        ):
            text = " ".join(header._element.xpath(".//w:t/text()"))
            self.assertIn("Current Example Client Ltd.", text)
            self.assertIn("Current Example Development", text)
            self.assertIn("Example County, Alberta", text)
            self.assertIn("Almor Proposal No.: P026-999", text)
            self.assertNotIn("Prior", text)
            self.assertNotIn("P026-111", text)
            self.assertNotIn("P026-222", text)

    def test_standard_field_program_keeps_fixed_template_wording(self):
        with tempfile.TemporaryDirectory() as directory:
            base, _ = make_standard_templates(directory)
            draft = sample_draft().model_copy(update={
                "field_program_intro": "AI must not replace the fixed field-program wording.",
            })
            result, _ = self.build(base, draft=draft)
        text = "\n".join(p.text for p in result.paragraphs)
        fixed = (
            "Based on information supplied by the Client, Almor provides the following estimated cost "
            "to perform the geotechnical assessment, and any additional work directed by the Client "
            "will constitute a change in the scope of work. "
            "The geotechnical field program could be initiated within two weeks subject to rig availability "
            "and locates, subsequent to authorization. The estimated costs do not include environmental "
            "considerations of the development.  The field program will include geotechnical boreholes "
            "advanced using a track mounted solid stem auger drilling rig. It is our understanding access "
            "and clearing, if required, to the boreholes is suitable for a track mounted drilling rig and "
            "will be provided by the client as this work is outside the scope of our proposal.  The proposed "
            "scope of work for this project includes:"
        )
        self.assertIn(fixed, text)
        self.assertNotIn("AI must not replace", text)

    def test_missing_dedicated_template_uses_only_the_fixed_approved_reference(self):
        with tempfile.TemporaryDirectory() as directory:
            library = Path(directory)
            base, _ = make_standard_templates(library)
            base.rename(library / "P026-121 Approved example.docx")
            decoy = Document(library / "P026-121 Approved example.docx")
            next(p for p in decoy.paragraphs if p.text == "Approved standard terms.").text = "Unwanted historical terms."
            decoy.save(library / "P026-182 Different example.docx")
            with patch.object(document_builder, "STANDARD_TEMPLATE_PATH", library / "missing.docx"), \
                 patch.object(document_builder, "HISTORICAL_DIR", library):
                data, _ = document_builder.build_standard_docx(sample_facts(), sample_draft())
            text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
            self.assertIn("Approved standard terms.", text)
            self.assertNotIn("Unwanted historical terms.", text)

    def test_ambiguous_fixed_reference_requires_one_standard_template(self):
        with tempfile.TemporaryDirectory() as directory:
            library = Path(directory)
            base, _ = make_standard_templates(library)
            data = base.read_bytes()
            (library / "P026-121 First example.docx").write_bytes(data)
            (library / "P026-121 Second example.docx").write_bytes(data)
            with patch.object(document_builder, "STANDARD_TEMPLATE_PATH", library / "missing.docx"), \
                 patch.object(document_builder, "HISTORICAL_DIR", library):
                with self.assertRaisesRegex(ValueError, "[Oo]ne|[Ss]ingle"):
                    document_builder.build_standard_docx(sample_facts(), sample_draft())

    def test_first_page_letterhead_keeps_its_absolute_drawing_position(self):
        with tempfile.TemporaryDirectory() as directory:
            base, _ = make_standard_templates(directory)
            source = Document(base)
            source.sections[0].header_distance = Inches(1)
            first_header = source.sections[0].first_page_header
            for child in list(first_header._element):
                first_header._element.remove(child)
            first_header.add_paragraph("Almor Company Letterhead")
            first_header._element.append(parse_xml(
                '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
                '<w:r><w:drawing><wp:anchor><wp:positionV relativeFrom="paragraph">'
                '<wp:posOffset>-91440</wp:posOffset></wp:positionV>'
                '</wp:anchor></w:drawing></w:r></w:p>'
            ))
            source.save(base)
            result, _ = self.build(base)
        first_header = result.sections[0].first_page_header
        self.assertIn("Almor Company Letterhead", " ".join(first_header._element.xpath(".//w:t/text()")))
        self.assertNotIn("Current Example Client Ltd.", " ".join(first_header._element.xpath(".//w:t/text()")))
        self.assertEqual(first_header._element.xpath(".//wp:positionV/@relativeFrom"), ["page"])
        self.assertEqual(first_header._element.xpath(".//wp:positionV/wp:posOffset/text()"), ["822960"])

    def test_cost_layout_and_standard_copy_are_fixed_across_different_projects(self):
        with tempfile.TemporaryDirectory() as directory:
            base, costs = make_standard_templates(directory)
            costs.unlink()  # An encrypted deployment need not have a separate cost-template DOCX.
            expected_widths = [Twips(value) for value in (3114, 3580, 528, 510, 1312, 928)]
            for project in ("Current Example Development", "A different project with different scope"):
                with self.subTest(project=project):
                    facts = sample_facts().model_copy(update={"project_name": project})
                    result, _ = self.build(base, facts)
                    self.assertEqual([column.width for column in result.tables[0].columns], expected_widths)
                    self.assertEqual([cell.text for cell in result.tables[0].rows[0].cells],
                                     ["Item", "Description", "Unit", "Est.", "Rate", "Total"])
                    item_paragraph = result.tables[0].rows[2].cells[0].paragraphs[0]
                    self.assertEqual(item_paragraph.runs[0].font.size, Pt(8))
                    self.assertEqual(item_paragraph.runs[0].font.name, "Arial")
                    self.assertEqual(item_paragraph.paragraph_format.space_after, Pt(8))
                    self.assertAlmostEqual(item_paragraph.paragraph_format.line_spacing, 1.07917, places=4)
                    self.assertFalse(result.tables[0].autofit)
                    self.assertFalse(result.tables[0]._tbl.xpath(".//w:gridSpan|.//w:vMerge"))
                    self.assertTrue(result.tables[0].rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
                    for row in result.tables[0].rows:
                        self.assertTrue(row._tr.xpath("./w:trPr/w:cantSplit"))
                    text = "\n".join(paragraph.text for paragraph in result.paragraphs)
                    self.assertIn("Approved standard terms.", text)
                    self.assertIn(
                        "Almor is interested in supporting the successful design and construction "
                        "of the proposed project.",
                        text,
                    )
                    self.assertNotIn("Approved closure one.", text)
                    self.assertNotIn("AI text must not replace", text)
                    table_text = "\n".join(cell.text for row in result.tables[0].rows for cell in row.cells)
                    self.assertIn("$450.00", table_text)
                    self.assertIn("The total estimated services cost is $450.00, excluding GST.", text)

    def test_ai_markdown_cost_table_cannot_create_a_second_cost_table(self):
        draft = sample_draft().model_copy(update={"cost_intro": (
            "The following estimate applies to this scope.\n"
            "Item | Description | Unit | Est. | Rate | Total\n"
            "Excavator | Field investigation | hr | 3 | $150.00 | $450.00"
        )})
        with tempfile.TemporaryDirectory() as directory:
            base, _ = make_standard_templates(directory)
            result, _ = self.build(base, draft=draft)
        paragraphs = [paragraph.text for paragraph in result.paragraphs]
        self.assertEqual(len(result.tables), 1)
        self.assertFalse(any("|" in paragraph for paragraph in paragraphs))
        self.assertIn("The following estimate applies to this scope.", paragraphs)

    def test_cost_item_rows_use_alternating_standard_shading(self):
        facts = sample_facts()
        facts.cost_items.append(CostLineItem(
            section="Field Program", item="Field supplies", description="Consumables",
            unit="LS", estimate=1, rate=50, total=50,
        ))
        with tempfile.TemporaryDirectory() as directory:
            base, _ = make_standard_templates(directory)
            result, _ = self.build(base, facts)
        rows = result.tables[0].rows
        fills = [row.cells[0]._tc.xpath("./w:tcPr/w:shd/@w:fill") for row in rows[2:4]]
        self.assertEqual(len(fills), 2)
        self.assertNotEqual(fills[0], fills[1])
        self.assertIn(["D6DBD9"], fills)

    def test_signature_group_stays_together_with_standard_font_sizes(self):
        with tempfile.TemporaryDirectory() as directory:
            base, _ = make_standard_templates(directory)
            result, _ = self.build(base)
        paragraphs = result.paragraphs
        label_index = next(
            index for index, paragraph in enumerate(paragraphs)
            if paragraph.text.startswith("Prepared by")
        )
        signature = paragraphs[label_index:]
        self.assertEqual(len(signature), 4)
        for index, paragraph in enumerate(signature):
            identifier = index == len(signature) - 1
            self.assertEqual(paragraph.paragraph_format.keep_with_next, not identifier)
            self.assertTrue(paragraph.runs)
            for run in paragraph.runs:
                self.assertEqual(run.font.name, "Arial")
                self.assertEqual(run.font.size, Pt(8 if identifier else 11))

    def test_cost_headers_and_section_labels_stay_with_the_first_item(self):
        with tempfile.TemporaryDirectory() as directory:
            base, _ = make_standard_templates(directory)
            result, _ = self.build(base)
        rows = result.tables[0].rows
        for row in rows[:2]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    self.assertTrue(paragraph.paragraph_format.keep_with_next)
        for cell in rows[2].cells:
            for paragraph in cell.paragraphs:
                self.assertFalse(paragraph.paragraph_format.keep_with_next)

    def test_closure_heading_and_signoff_stay_with_the_signature_block(self):
        with tempfile.TemporaryDirectory() as directory:
            base, _ = make_standard_templates(directory)
            result, _ = self.build(base)
        paragraphs = result.paragraphs
        start = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text == "Closure")
        closing = paragraphs[start:]
        for paragraph in closing[:-1]:
            self.assertTrue(paragraph.paragraph_format.keep_with_next)
        self.assertFalse(closing[-1].paragraph_format.keep_with_next)


if __name__ == "__main__":
    unittest.main()
