import io
import unittest

from docx import Document

from proposal_app.uploaded_proposal import extract_uploaded_proposal


def reviewed_proposal(client_line="", introduction=""):
    document = Document()
    document.add_paragraph("October 28, 2026\tAlmor Proposal No.: P026-200")
    if client_line:
        document.add_paragraph(client_line)
    document.add_paragraph("Attention: Ms. Alex Example | alex@example.com\t")
    document.add_paragraph("Re:\tRequest for Proposal for Geotechnical Services\nRoad Assessment\nCalgary")
    document.add_heading("Introduction", 1)
    document.add_paragraph(introduction)
    document.add_paragraph("The total estimated cost is $2,221.00.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class UploadedProposalFieldsTest(unittest.TestCase):
    def test_reads_individual_client_from_introduction_without_client_block(self):
        data = reviewed_proposal(introduction=(
            "At the request of Alex Example, Almor Testing Services Ltd. (Almor) is "
            "pleased to provide this proposal for geotechnical services."
        ))
        result = extract_uploaded_proposal(data, "P026-200 Road Assessment.docx")
        self.assertEqual(result.facts.client_name, "Alex Example")
        self.assertEqual(result.facts.contact_email, "alex@example.com")
        self.assertEqual(result.facts.proposal_number, "P026-200")
        self.assertEqual(result.facts.project_name, "Road Assessment")
        self.assertEqual(result.budget, 2221)

    def test_prefers_explicit_client_block_over_introduction(self):
        data = reviewed_proposal("Example Engineering Ltd.\n123 Test Avenue", (
            "At the request of Alex Example, Almor Testing Services Ltd. is pleased to provide this proposal."
        ))
        result = extract_uploaded_proposal(data)
        self.assertEqual(result.facts.client_name, "Example Engineering Ltd.")
        self.assertEqual(result.facts.client_address, "123 Test Avenue")

    def test_reads_company_from_requester_in_introduction(self):
        data = reviewed_proposal(introduction=(
            "At the request of Alex Example of Example Engineering Ltd. (Client), "
            "Almor Testing Services Ltd. (Almor) is pleased to provide this proposal."
        ))
        result = extract_uploaded_proposal(data)
        self.assertEqual(result.facts.client_name, "Example Engineering Ltd.")

    def test_does_not_assume_attention_contact_is_the_client(self):
        data = reviewed_proposal(introduction="Almor is pleased to provide this proposal.")
        with self.assertRaisesRegex(ValueError, "required proposal fields: client"):
            extract_uploaded_proposal(data)

    def test_preserves_organization_name_containing_of(self):
        data = reviewed_proposal(introduction=(
            "At the request of City of Calgary (Client), Almor Testing Services Ltd. "
            "is pleased to provide this proposal."
        ))
        self.assertEqual(extract_uploaded_proposal(data).facts.client_name, "City of Calgary")

    def test_preserves_company_punctuation(self):
        data = reviewed_proposal(introduction=(
            "At the request of Example Engineering, Inc. (Client), Almor Testing Services Ltd. "
            "is pleased to provide this proposal."
        ))
        self.assertEqual(extract_uploaded_proposal(data).facts.client_name, "Example Engineering, Inc.")
